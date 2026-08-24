"""
src/infrastructure/loaders/pptx_loader.py

IDocumentLoader adapter backed by python-pptx.

One Document per slide (page = slide number) — slides are already a
natural, bounded unit of meaning, unlike DOCX's continuous prose which
needs char-count pseudo-page grouping.

Extracts, in shape order: text frames (titles, bullets, body text) and
tables (converted to Markdown — same convention as DocxDocumentLoader/
PdfDocumentLoader, so MetadataEnricher's has_tables detection and the
table-aware RecursiveTextChunker both pick these up automatically, no
changes needed there). Speaker notes, when present, are appended after
the visible slide content under a "--- Speaker Notes ---" separator,
since they often carry the explanatory detail behind terse bullet points.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.table import Table

from src.domain.entities import Document
from src.domain.interfaces import IDocumentLoader, ILogger

SUPPORTED_EXTENSION = ".pptx"
_NOTES_SEPARATOR = "--- Speaker Notes ---"


class PptxDocumentLoader(IDocumentLoader):
    def __init__(self, logger: ILogger) -> None:
        self._logger = logger

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == SUPPORTED_EXTENSION

    def load(self, path: Path) -> list[Document]:
        presentation = Presentation(str(path))
        total = len(presentation.slides)
        docs: list[Document] = []
        slides_with_tables = 0
        slides_with_notes = 0

        for i, slide in enumerate(presentation.slides):
            blocks, table_count = self._extract_slide_blocks(slide)
            notes = self._extract_notes(slide)

            if not blocks and not notes:
                self._logger.warning(
                    f"Slide {i + 1}/{total} of '{path.name}' yielded no content — skipping."
                )
                continue

            content = "\n\n".join(blocks)
            if notes:
                content = f"{content}\n\n{_NOTES_SEPARATOR}\n{notes}" if content else notes
                slides_with_notes += 1
            if table_count:
                slides_with_tables += 1

            docs.append(Document(
                page_content=content,
                metadata={
                    "source": path.name,
                    "source_path": str(path),
                    "page": i + 1,
                    "total_pages": total,
                    "file_type": "pptx",
                    "table_count": table_count,
                    "has_notes": bool(notes),
                },
            ))

        self._logger.info(
            f"Loaded '{path.name}': {len(docs)}/{total} slide(s) with content "
            f"({slides_with_tables} with tables, {slides_with_notes} with speaker notes)."
        )
        return docs

    def _extract_slide_blocks(self, slide) -> tuple[list[str], int]:
        """
        Extracts text frames and tables from a slide's shapes, in shape
        order. Returns (blocks, table_count).
        """
        blocks: list[str] = []
        table_count = 0

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    blocks.append(text)

            if shape.has_table:
                md = self._table_to_markdown(shape.table)
                if md:
                    blocks.append(md)
                    table_count += 1

        return blocks, table_count

    @staticmethod
    def _extract_notes(slide) -> str:
        if not slide.has_notes_slide:
            return ""
        notes_frame = slide.notes_slide.notes_text_frame
        return (notes_frame.text or "").strip()

    @staticmethod
    def _table_to_markdown(table: Table) -> str:
        """Convert a python-pptx Table to a Markdown table string."""
        rows: list[str] = []
        for i, row in enumerate(table.rows):
            cells = [
                cell.text.strip().replace("|", "\\|").replace("\n", " ")
                for cell in row.cells
            ]
            if not any(cells):
                continue
            rows.append("| " + " | ".join(cells) + " |")
            if i == 0:
                rows.append("| " + " | ".join(["---"] * len(cells)) + " |")

        return "\n".join(rows)